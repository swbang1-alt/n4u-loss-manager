엔포유 소분 작업 및 감모 관리 업로드 프로그램.

실행:
streamlit run n4u_upload_app.py
"""

from future import annotations

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any, BinaryIO

import openpyxl
import pandas as pd
import streamlit as st

=========================================================

화면 표시 컬럼

=========================================================

DISPLAY_COLUMNS = [
"거래일자",
"전표번호",
"투입품목코드",
"투입품목명",
"투입수량",
"투입중량(kg)",
"투입원가",
"투입원가(합계)",
"산출품목코드",
"산출품목명",
"제조수량",
"제조중량(kg)",
"납품원가",
"납품원가(합계)",
"감모수량(kg)",
"감모금액(합계)",
"감모비율(%)",
"작업비비율(%)",
"소분작업비용(합계)",
]

=========================================================

제품코드 정리

=========================================================

def normalize_code(value: Any) -> str:
"""엑셀의 숫자/문자 제품코드를 같은 문자열로 맞춘다."""

if value is None:
    return ""

if isinstance(value, float) and value.is_integer():
    return str(int(value))

text = str(value).strip()

if re.fullmatch(r"\d+\.0", text):
    return text[:-2]

return text

=========================================================

숫자 변환

=========================================================

def to_number(
value: Any,
default: float = 0.0,
) -> float:

if value is None:
    return default

if str(value).strip() == "":
    return default

if isinstance(value, (int, float)):
    return float(value)

cleaned = re.sub(
    r"[^0-9.+-]",
    "",
    str(value).replace(",", ""),
)

try:
    return float(cleaned)
except ValueError:
    return default

=========================================================

엑셀 컬럼 찾기

=========================================================

def find_columns(
rows: list[tuple[Any, ...]],
) -> tuple[int, dict[str, int]]:
"""소분리스트 첫 15행에서 보고서 헤더와 필요한 열을 찾는다."""

patterns = {
    "date": [
        "거래년월일",
        "거래일자",
        "일자",
    ],
    "doc": [
        "전표번호",
        "전표no",
        "전표",
    ],
    "warehouse": [
        "창고명",
        "창고",
    ],
    "qty": [
        "수량",
    ],
    "name": [
        "상품명",
        "품명",
        "내역",
        "품목명",
    ],
    "code": [
        "품목코드",
        "상품코드",
        "코드",
    ],
}

for row_index, row in enumerate(rows[:15]):

    found: dict[str, int] = {}

    for col_index, cell in enumerate(row):

        header = str(
            cell or ""
        ).replace(
            " ",
            "",
        ).lower()

        for key, words in patterns.items():

            if (
                key not in found
                and any(
                    word in header
                    for word in words
                )
            ):
                found[key] = col_index

        # 판매금액 우선
        if header == "판매금액":

            found["price"] = col_index

        # 판매금액이 없는 기존 양식
        elif (
            header == "합계금액"
            and "price" not in found
        ):

            found["price"] = col_index

        elif (
            header == "공급가액"
            and "price" not in found
        ):

            found["price"] = col_index

    if (
        "warehouse" in found
        and "price" in found
    ):
        return row_index, found

raise ValueError(
    "소분리스트에서 필수 열인 "
    "'창고명'과 '판매금액/금액'을 찾지 못했습니다."
)

=========================================================

날짜 추출

=========================================================

def extract_date(
value: Any,
) -> str:

if isinstance(
    value,
    (datetime, date),
):
    return value.strftime(
        "%Y-%m-%d"
    )

match = re.search(
    r"20\d{2}[.\-/]\d{1,2}[.\-/]\d{1,2}",
    str(value or ""),
)

if match:

    return (
        match.group()
        .replace(".", "-")
        .replace("/", "-")
    )

return ""

=========================================================

특정 값 검색

=========================================================

def first_match(
row: tuple[Any, ...],
pattern: str,
) -> str:

for cell in row:

    match = re.search(
        pattern,
        str(cell or ""),
    )

    if match:
        return match.group()

return ""

=========================================================

제품 중량 읽기

=========================================================

def read_product_weights(
file: BinaryIO,
) -> dict[str, float]:

workbook = openpyxl.load_workbook(
    file,
    data_only=True,
    read_only=True,
)

sheet = workbook.active

weights: dict[str, float] = {}

for row in sheet.iter_rows(
    values_only=True
):

    if len(row) <= 26:
        continue

    # B열 : 제품코드
    code = normalize_code(
        row[1]
    )

    # AA열 : 중량
    weight = to_number(
        row[26],
        default=-1,
    )

    if (
        code
        and weight >= 0
    ):
        weights[code] = weight

return weights

=========================================================

파일 분석

=========================================================

def process_files(
list_file: BinaryIO,
product_file: BinaryIO,
) -> list[dict[str, Any]]:

# -----------------------------------------------------
# 제품 중량 불러오기
# -----------------------------------------------------
weights = read_product_weights(
    product_file
)

# -----------------------------------------------------
# 소분리스트 불러오기
# -----------------------------------------------------
workbook = openpyxl.load_workbook(
    list_file,
    data_only=True,
    read_only=True,
)

rows = list(
    workbook.active.iter_rows(
        values_only=True
    )
)

# -----------------------------------------------------
# 헤더 찾기
# -----------------------------------------------------
header_index, columns = find_columns(
    rows
)

# -----------------------------------------------------
# 전표별 그룹
# -----------------------------------------------------
grouped: dict[
    tuple[str, str],
    dict[str, Any],
] = {}

current_doc = "-"
current_date = "-"

# =====================================================
# 데이터 행 처리
# =====================================================
for row in rows[
    header_index + 1:
]:

    if (
        not row
        or len(row)
        <= max(
            columns["warehouse"],
            columns["price"],
        )
    ):
        continue

    # -------------------------------------------------
    # 합계/소계 행 제외
    # -------------------------------------------------
    row_text = " ".join(
        str(cell)
        for cell in row
        if cell is not None
    )

    if any(
        word in row_text
        for word in (
            "합계",
            "합 계",
            "소계",
            "소 계",
        )
    ):
        continue

    # -------------------------------------------------
    # 창고
    # -------------------------------------------------
    warehouse = str(
        row[
            columns["warehouse"]
        ]
        or ""
    ).strip()

    # -------------------------------------------------
    # 금액
    # -------------------------------------------------
    raw_price = row[
        columns["price"]
    ]

    if (
        not warehouse
        or raw_price is None
    ):
        continue

    # -------------------------------------------------
    # 전표번호
    # -------------------------------------------------
    if "doc" in columns:

        doc = first_match(
            (
                row[
                    columns["doc"]
                ],
            ),
            r"\d{6,}",
        )

    else:

        doc = ""

    doc = (
        doc
        or first_match(
            row,
            r"\d{6,}",
        )
        or current_doc
    )

    current_doc = doc

    # -------------------------------------------------
    # 거래일자
    # -------------------------------------------------
    if "date" in columns:

        transaction_date = extract_date(
            row[
                columns["date"]
            ]
        )

    else:

        transaction_date = ""

    transaction_date = (
        transaction_date
        or first_match(
            row,
            r"20\d{2}[.\-/]\d{1,2}[.\-/]\d{1,2}",
        )
    )

    if transaction_date:

        transaction_date = (
            transaction_date
            .replace(".", "-")
            .replace("/", "-")
        )

    else:

        transaction_date = current_date

    current_date = transaction_date

    if transaction_date == "-":
        continue

    # -------------------------------------------------
    # 수량
    # -------------------------------------------------
    if "qty" in columns:

        qty = to_number(
            row[
                columns["qty"]
            ],
            1.0,
        )

    else:

        qty = 1.0

    # -------------------------------------------------
    # 품목코드
    # -------------------------------------------------
    if "code" in columns:

        item_code = normalize_code(
            row[
                columns["code"]
            ]
        )

    else:

        item_code = ""

    # -------------------------------------------------
    # 품목명
    # -------------------------------------------------
    if "name" in columns:

        item_name = str(
            row[
                columns["name"]
            ]
            or ""
        ).strip()

    else:

        item_name = ""

    # -------------------------------------------------
    # 단중량
    # -------------------------------------------------
    unit_weight = weights.get(
        item_code,
        0.0,
    )

    # -------------------------------------------------
    # 상세 데이터
    # -------------------------------------------------
    detail = {
        "code": item_code or "-",
        "name": item_name or "-",
        "qty": qty,
        "price": to_number(
            raw_price
        ),
        "unit_weight": unit_weight,
        "weight_kg": (
            qty
            * unit_weight
            / 1000
        ),
    }

    # -------------------------------------------------
    # 전표 그룹 생성
    # -------------------------------------------------
    record = grouped.setdefault(
        (
            doc,
            transaction_date,
        ),
        {
            "inp": [],
            "out": [],
        },
    )

    # -------------------------------------------------
    # 소분 창고 = 투입
    # -------------------------------------------------
    if "소분" in warehouse:

        record[
            "inp"
        ].append(
            detail
        )

    # -------------------------------------------------
    # 대연점 창고 = 산출
    # -------------------------------------------------
    elif "대연점" in warehouse:

        record[
            "out"
        ].append(
            detail
        )

# =====================================================
# 전표별 계산
# =====================================================
result: list[
    dict[str, Any]
] = []

for (
    doc,
    transaction_date,
), record in sorted(
    grouped.items(),
    key=lambda item: (
        item[0][1],
        item[0][0],
    ),
):

    # -------------------------------------------------
    # 합계 함수
    # -------------------------------------------------
    def total(
        side: str,
        key: str,
    ) -> float:

        return sum(
            float(item[key])
            for item in record[
                side
            ]
        )

    # =================================================
    # 투입
    # =================================================
    inp_qty = total(
        "inp",
        "qty",
    )

    inp_cost = total(
        "inp",
        "price",
    )

    inp_weight = total(
        "inp",
        "weight_kg",
    )

    # =================================================
    # 산출
    # =================================================
    out_qty = total(
        "out",
        "qty",
    )

    out_cost = total(
        "out",
        "price",
    )

    out_weight = total(
        "out",
        "weight_kg",
    )

    # =================================================
    # 투입원가
    # 수량 1개당 투입원가
    # =================================================
    inp_unit_cost = (
        abs(inp_cost)
        / abs(inp_qty)
        if inp_qty
        else 0.0
    )

    # =================================================
    # 납품원가
    # 수량 1개당 납품원가
    # =================================================
    out_unit_cost = (
        abs(out_cost)
        / abs(out_qty)
        if out_qty
        else 0.0
    )

    # =================================================
    # 감모수량
    # =================================================
    loss_weight = (
        inp_weight
        + out_weight
    )

    # =================================================
    # 소분작업비용
    # =================================================
    work_cost = (
        inp_cost
        + out_cost
    )

    # =================================================
    # 감모금액
    #
    # 투입 kg당 원가 × 감모중량
    # =================================================
    loss_unit_cost = (
        abs(inp_cost)
        / abs(inp_weight)
        if inp_weight
        else 0.0
    )

    loss_cost = (
        loss_weight
        * loss_unit_cost
    )

    # =================================================
    # 감모비율
    # =================================================
    loss_ratio = (
        loss_weight
        / abs(inp_weight)
        * 100
        if inp_weight
        else 0.0
    )

    # =================================================
    # 작업비비율
    # =================================================
    work_ratio = (
        work_cost
        / -inp_cost
        * 100
        if inp_cost
        else 0.0
    )

    # =================================================
    # 결과 저장
    # =================================================
    result.append(
        {
            "date": transaction_date,
            "doc_no": doc,

            # 투입
            "inp_code": ", ".join(
                sorted(
                    {
                        item["code"]
                        for item
                        in record["inp"]
                        if item["code"] != "-"
                    }
                )
            ) or "-",

            "inp_name": ", ".join(
                sorted(
                    {
                        item["name"]
                        for item
                        in record["inp"]
                        if item["name"] != "-"
                    }
                )
            ) or "-",

            "inp_qty": inp_qty,
            "inp_cost_unit": inp_unit_cost,
            "inp_cost": inp_cost,
            "inp_wt_kg": inp_weight,

            # 산출
            "out_code": ", ".join(
                sorted(
                    {
                        item["code"]
                        for item
                        in record["out"]
                        if item["code"] != "-"
                    }
                )
            ) or "-",

            "out_name": ", ".join(
                sorted(
                    {
                        item["name"]
                        for item
                        in record["out"]
                        if item["name"] != "-"
                    }
                )
            ) or "-",

            "out_qty": out_qty,
            "out_cost_unit": out_unit_cost,
            "out_cost": out_cost,
            "out_wt_kg": out_weight,

            # 감모
            "loss_wt_kg": loss_weight,
            "loss_cost": loss_cost,

            # 비율
            "loss_ratio": loss_ratio,
            "ratio": work_ratio,

            # 작업비
            "work_cost": work_cost,

            # 상세
            "inp_details": record[
                "inp"
            ],

            "out_details": record[
                "out"
            ],
        }
    )

return result

=========================================================

화면용 DataFrame

=========================================================

def display_dataframe(
records: list[dict[str, Any]],
) -> pd.DataFrame:

rows = []

for item in records:

    rows.append(
        {
            "거래일자": item[
                "date"
            ],

            "전표번호": item[
                "doc_no"
            ],

            "투입품목코드": item[
                "inp_code"
            ],

            "투입품목명": item[
                "inp_name"
            ],

            "투입수량": item[
                "inp_qty"
            ],

            "투입중량(kg)": item[
                "inp_wt_kg"
            ],

            "투입원가": item[
                "inp_cost_unit"
            ],

            "투입원가(합계)": item[
                "inp_cost"
            ],

            "산출품목코드": item[
                "out_code"
            ],

            "산출품목명": item[
                "out_name"
            ],

            "제조수량": item[
                "out_qty"
            ],

            "제조중량(kg)": item[
                "out_wt_kg"
            ],

            "납품원가": item[
                "out_cost_unit"
            ],

            "납품원가(합계)": item[
                "out_cost"
            ],

            "감모수량(kg)": item[
                "loss_wt_kg"
            ],

            "감모금액(합계)": item[
                "loss_cost"
            ],

            "감모비율(%)": item[
                "loss_ratio"
            ],

            "작업비비율(%)": item[
                "ratio"
            ],

            "소분작업비용(합계)": item[
                "work_cost"
            ],
        }
    )

return pd.DataFrame(
    rows,
    columns=DISPLAY_COLUMNS,
)

=========================================================

일자별 소계 / 전체 총합계

=========================================================

def with_daily_subtotals(
dataframe: pd.DataFrame,
) -> pd.DataFrame:

"""전표 목록에 일자별 소계와 전체 총합계 행을 덧붙인다."""

numeric_columns = [
    "투입수량",
    "투입중량(kg)",
    "투입원가(합계)",
    "제조수량",
    "제조중량(kg)",
    "납품원가(합계)",
    "감모수량(kg)",
    "감모금액(합계)",
    "소분작업비용(합계)",
]

# =====================================================
# 소계 행 생성
# =====================================================
def summary_row(
    rows: pd.DataFrame,
    label: str,
) -> dict[str, Any]:

    totals = rows[
        numeric_columns
    ].sum()

    # -------------------------------------------------
    # 투입
    # -------------------------------------------------
    input_qty = totals[
        "투입수량"
    ]

    input_weight = totals[
        "투입중량(kg)"
    ]

    input_cost = totals[
        "투입원가(합계)"
    ]

    # -------------------------------------------------
    # 납품
    # -------------------------------------------------
    output_qty = totals[
        "제조수량"
    ]

    output_cost = totals[
        "납품원가(합계)"
    ]

    # -------------------------------------------------
    # 감모
    # -------------------------------------------------
    loss_weight = totals[
        "감모수량(kg)"
    ]

    loss_cost = totals[
        "감모금액(합계)"
    ]

    # -------------------------------------------------
    # 투입원가 낱개
    # -------------------------------------------------
    input_unit_cost = (
        abs(input_cost)
        / abs(input_qty)
        if input_qty
        else 0.0
    )

    # -------------------------------------------------
    # 납품원가 낱개
    # -------------------------------------------------
    output_unit_cost = (
        abs(output_cost)
        / abs(output_qty)
        if output_qty
        else 0.0
    )

    # -------------------------------------------------
    # 소계 결과
    # -------------------------------------------------
    return {
        "거래일자": label,

        "전표번호": "-",

        "투입품목코드": "-",
        "투입품목명": "-",

        "투입수량": input_qty,
        "투입중량(kg)": input_weight,

        "투입원가": input_unit_cost,
        "투입원가(합계)": input_cost,

        "산출품목코드": "-",
        "산출품목명": "-",

        "제조수량": output_qty,
        "제조중량(kg)": totals[
            "제조중량(kg)"
        ],

        "납품원가": output_unit_cost,
        "납품원가(합계)": output_cost,

        "감모수량(kg)": loss_weight,

        "감모금액(합계)": loss_cost,

        "감모비율(%)": (
            loss_weight
            / abs(input_weight)
            * 100
            if input_weight
            else 0.0
        ),

        "작업비비율(%)": (
            totals[
                "소분작업비용(합계)"
            ]
            / -input_cost
            * 100
            if input_cost
            else 0.0
        ),

        "소분작업비용(합계)": totals[
            "소분작업비용(합계)"
        ],
    }

# =====================================================
# 결과 조립
# =====================================================
result_parts: list[
    pd.DataFrame
] = []

for (
    transaction_date,
    rows,
) in dataframe.groupby(
    "거래일자",
    sort=True,
):

    # 실제 전표
    result_parts.append(
        rows
    )

    # 일자별 소계
    result_parts.append(
        pd.DataFrame(
            [
                summary_row(
                    rows,
                    f"[{transaction_date} 소계]",
                )
            ]
        )
    )

# 전체 총합계
result_parts.append(
    pd.DataFrame(
        [
            summary_row(
                dataframe,
                "★ 전체 총합계",
            )
        ]
    )
)

return pd.concat(
    result_parts,
    ignore_index=True,
)[DISPLAY_COLUMNS]

=========================================================

엑셀 다운로드

=========================================================

def to_excel(
dataframe: pd.DataFrame,
) -> bytes:

output = BytesIO()

with pd.ExcelWriter(
    output,
    engine="openpyxl",
) as writer:

    dataframe.to_excel(
        writer,
        index=False,
        sheet_name="소분작업분석",
    )

return output.getvalue()

=========================================================

Streamlit 기본 설정

=========================================================

st.set_page_config(
page_title="엔포유 소분 작업 관리",
page_icon="/scale:",
layout="wide",
)

=========================================================

제목

=========================================================

st.title(
"엔포유 소분 작업 및 감모 관리"
)

st.caption(
"소분리스트와 제품조회 파일을 업로드하면 "
"전표별 투입·산출·감모 현황을 계산합니다."
)

=========================================================

Session State

=========================================================

for key, default in {
"records": None,
"file_signature": None,
}.items():

st.session_state.setdefault(
    key,
    default,
)

=========================================================

파일 업로드

=========================================================

with st.container(
border=True
):

st.subheader(
    "파일 업로드",
    divider="gray",
)

upload_list, upload_product = st.columns(
    2
)

# -----------------------------------------------------
# 소분리스트
# -----------------------------------------------------
with upload_list:

    list_file = st.file_uploader(
        "1. 엔포유소분리스트",
        type=["xlsx"],
        key="list_file",
    )

# -----------------------------------------------------
# 제품조회
# -----------------------------------------------------
with upload_product:

    product_file = st.file_uploader(
        "2. 제품조회",
        type=["xlsx"],
        key="product_file",
    )

# =====================================================
# 파일이 모두 올라온 경우
# =====================================================
if (
    list_file
    and product_file
):

    signature = (
        list_file.name,
        list_file.size,
        product_file.name,
        product_file.size,
    )

    if (
        signature
        != st.session_state.file_signature
    ):

        try:

            with st.spinner(
                "업로드 파일을 분석하는 중입니다..."
            ):

                st.session_state.records = (
                    process_files(
                        BytesIO(
                            list_file.getvalue()
                        ),
                        BytesIO(
                            product_file.getvalue()
                        ),
                    )
                )

                st.session_state.file_signature = (
                    signature
                )

            st.success(
                f"분석 완료: 전표 "
                f"{len(st.session_state.records):,}건"
            )

        except Exception as exc:

            st.session_state.records = None

            st.error(
                "파일을 처리하지 못했습니다. "
                "엑셀 형식을 확인해 주세요.\n\n"
                f"{exc}"
            )

else:

    st.info(
        "두 개의 .xlsx 파일을 모두 올려 주세요."
    )

=========================================================

분석 결과

=========================================================

records = st.session_state.records

if records is not None:

# -----------------------------------------------------
# 분석 대상 없음
# -----------------------------------------------------
if not records:

    st.warning(
        "소분 또는 대연점 창고에 해당하는 "
        "분석 대상 행을 찾지 못했습니다."
    )

    st.stop()

# -----------------------------------------------------
# 화면 DataFrame
# -----------------------------------------------------
frame = display_dataframe(
    records
)

# -----------------------------------------------------
# 날짜 범위
# -----------------------------------------------------
available_dates = pd.to_datetime(
    frame["거래일자"],
    errors="coerce",
).dropna()

min_date = (
    available_dates
    .min()
    .date()
)

max_date = (
    available_dates
    .max()
    .date()
)


# =====================================================
# 조회 조건
# =====================================================
with st.container(
    border=True
):

    st.subheader(
        "조회 조건",
        divider="gray",
    )

    with st.form(
        "filter_form",
        border=False,
    ):

        start_column, end_column = st.columns(
            2
        )

        # -------------------------------------------------
        # 시작일
        # -------------------------------------------------
        with start_column:

            start_date = st.date_input(
                "시작일",
                value=min_date,
                key="filter_start_date",
            )

        # -------------------------------------------------
        # 종료일
        # -------------------------------------------------
        with end_column:

            end_date = st.date_input(
                "종료일",
                value=max_date,
                key="filter_end_date",
            )

        # -------------------------------------------------
        # 품목명
        # -------------------------------------------------
        keyword = st.text_input(
            "투입품목명",
            placeholder="예: 삼겹살",
        )

        # -------------------------------------------------
        # 조회 버튼
        # -------------------------------------------------
        st.form_submit_button(
            "조회",
            icon=":material/search:",
            type="primary",
        )


# =====================================================
# 날짜 오류
# =====================================================
if start_date > end_date:

    st.error(
        "시작일은 종료일보다 늦을 수 없습니다."
    )

    st.stop()


# =====================================================
# 데이터 필터
# =====================================================
filtered = frame[
    (
        frame["거래일자"]
        >= start_date.isoformat()
    )
    & (
        frame["거래일자"]
        <= end_date.isoformat()
    )
]

# -----------------------------------------------------
# 품목명 필터
# -----------------------------------------------------
if keyword:

    filtered = filtered[
        filtered[
            "투입품목명"
        ].str.contains(
            keyword,
            case=False,
            na=False,
        )
    ]


# =====================================================
# KPI 계산
# =====================================================
total_input_weight = filtered[
    "투입중량(kg)"
].sum()

total_loss_weight = filtered[
    "감모수량(kg)"
].sum()

total_input_cost = filtered[
    "투입원가(합계)"
].sum()

total_work_cost = filtered[
    "소분작업비용(합계)"
].sum()

total_loss_ratio = (
    total_loss_weight
    / abs(total_input_weight)
    * 100
    if total_input_weight
    else 0.0
)

total_work_ratio = (
    total_work_cost
    / -total_input_cost
    * 100
    if total_input_cost
    else 0.0
)


# =====================================================
# KPI 표시
# =====================================================
metrics = st.columns(
    7,
    wrap=False,
)

metrics[0].metric(
    "조회 전표",
    f"{len(filtered):,}건",
)

metrics[1].metric(
    "총 투입중량",
    f"{total_input_weight:,.2f} kg",
)

metrics[2].metric(
    "총 감모중량",
    f"{total_loss_weight:,.2f} kg",
)

metrics[3].metric(
    "총 감모금액",
    f"{filtered['감모금액(합계)'].sum():,.0f}원",
)

metrics[4].metric(
    "총 감모비율",
    f"{total_loss_ratio:,.2f}%",
)

metrics[5].metric(
    "총 소분작업비용",
    f"{total_work_cost:,.0f}원",
)

metrics[6].metric(
    "총 작업비비율",
    f"{total_work_ratio:,.2f}%",
)


# =====================================================
# 전표별 분석 결과
# =====================================================
st.subheader(
    "전표별 분석 결과"
)

table_frame = with_daily_subtotals(
    filtered
)


# =====================================================
# 결과 테이블
# =====================================================
st.dataframe(
    table_frame,
    hide_index=True,

    column_config={

        # -------------------------------------------------
        # 중량
        # -------------------------------------------------
        name: st.column_config.NumberColumn(
            format="%,.2f"
        )

        for name in [
            "투입중량(kg)",
            "제조중량(kg)",
            "감모수량(kg)",
        ]
    }

    |

    {

        # -------------------------------------------------
        # 금액 / 수량
        # -------------------------------------------------
        name: st.column_config.NumberColumn(
            format="%,.0f"
        )

        for name in [
            "투입수량",
            "투입원가",
            "투입원가(합계)",
            "제조수량",
            "납품원가",
            "납품원가(합계)",
            "감모금액(합계)",
            "소분작업비용(합계)",
        ]
    }

    |

    {

        # -------------------------------------------------
        # 비율
        # -------------------------------------------------
        name: st.column_config.NumberColumn(
            format="%.2f%%"
        )

        for name in [
            "감모비율(%)",
            "작업비비율(%)",
        ]
    },

    key="results_table",
)


# =====================================================
# 엑셀 다운로드
# =====================================================
st.download_button(
    "결과 엑셀 다운로드",

    data=to_excel(
        table_frame
    ),

    file_name=(
        "엔포유_소분작업분석.xlsx"
    ),

    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),

    icon=":material/download:",
)


# =====================================================
# 전표 상세 조회
# =====================================================
lookup = {
    (
        item["date"],
        item["doc_no"],
    ): item

    for item in records
}

choices = [
    (
        row["거래일자"],
        row["전표번호"],
    )

    for _, row in filtered.iterrows()
]


# =====================================================
# 상세 전표가 있는 경우
# =====================================================
if choices:

    st.subheader(
        "전표 상세"
    )

    selected = st.selectbox(
        "상세 전표",
        choices,

        format_func=lambda value:
            f"{value[0]} | 전표번호 {value[1]}",
    )

    detail = lookup[
        selected
    ]


    # =================================================
    # 상세 원가
    # =================================================
    cost_metrics = st.columns(
        5
    )

    # -------------------------------------------------
    # 투입원가
    # -------------------------------------------------
    cost_metrics[0].metric(
        "투입원가",
        f"{detail['inp_cost_unit']:,.0f}원",
    )

    # -------------------------------------------------
    # 투입원가 합계
    # -------------------------------------------------
    cost_metrics[1].metric(
        "투입원가(합계)",
        f"{detail['inp_cost']:,.0f}원",
    )

    # -------------------------------------------------
    # 납품원가
    # -------------------------------------------------
    cost_metrics[2].metric(
        "납품원가",
        f"{detail['out_cost_unit']:,.0f}원",
    )

    # -------------------------------------------------
    # 납품원가 합계
    # -------------------------------------------------
    cost_metrics[3].metric(
        "납품원가(합계)",
        f"{detail['out_cost']:,.0f}원",
    )

    # -------------------------------------------------
    # 소분작업비용
    # -------------------------------------------------
    cost_metrics[4].metric(
        "소분작업비용",
        f"{detail['work_cost']:,.0f}원",
    )


    # =================================================
    # 투입 / 산출 상세
    # =================================================
    left, right = st.columns(
        2
    )

    for (
        column,
        title,
        items,
    ) in (
        (
            left,
            "투입 품목 (소분)",
            detail["inp_details"],
        ),
        (
            right,
            "산출 품목 (대연점)",
            detail["out_details"],
        ),
    ):

        with column.container(
            border=True
        ):

            st.markdown(
                f"#### {title}"
            )

            detail_frame = (
                pd.DataFrame(
                    items
                )
                .rename(
                    columns={
                        "code": "코드",
                        "name": "품목명",
                        "qty": "수량",
                        "price": "금액",
                        "unit_weight": "단중량(g)",
                        "weight_kg": "총중량(kg)",
                    }
                )
            )

            st.dataframe(
                detail_frame,
                hide_index=True,
            )
